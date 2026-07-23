from __future__ import annotations

import json
import resource
import shutil
import subprocess
import sys
from pathlib import Path


def apply_limits() -> None:
    try:
        resource.setrlimit(resource.RLIMIT_CPU, (12, 12))
        memory = 512 * 1024 * 1024
        resource.setrlimit(resource.RLIMIT_AS, (memory, memory))
        resource.setrlimit(resource.RLIMIT_FSIZE, (20 * 1024 * 1024, 20 * 1024 * 1024))
    except (ValueError, OSError):
        pass


def extract(path: Path, kind: str) -> str:
    if kind == "pdf":
        pdftotext = shutil.which("pdftotext")
        if pdftotext:
            result = subprocess.run(
                [pdftotext, "-layout", "-nopgbrk", str(path), "-"],
                capture_output=True,
                text=True,
                timeout=15,
                check=False,
            )
            if result.returncode == 0 and len(result.stdout.strip()) >= 40:
                return result.stdout
        from pypdf import PdfReader

        reader = PdfReader(str(path), strict=False)
        if len(reader.pages) > 20:
            raise ValueError("Resume PDF exceeds the 20-page safety limit")
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if kind == "docx":
        from docx import Document

        document = Document(str(path))
        lines = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)

    return path.read_text(encoding="utf-8", errors="replace")


def main() -> int:
    apply_limits()
    if len(sys.argv) != 3:
        print(json.dumps({"error": "expected path and kind"}))
        return 2
    try:
        text = extract(Path(sys.argv[1]), sys.argv[2])
        print(json.dumps({"text": text}, ensure_ascii=False))
        return 0
    except Exception as exc:
        print(json.dumps({"error": str(exc)}))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
